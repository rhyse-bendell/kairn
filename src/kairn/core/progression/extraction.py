from __future__ import annotations
import hashlib, html.parser, re
from pathlib import Path
from kairn.core.storage.sqlite import connect, init_db

def normalize_text(s: str) -> str:
    return re.sub(r"\s+", " ", (s or '').strip().lower())
def content_hash(text: str) -> str: return hashlib.sha1((text or '').encode('utf-8')).hexdigest()
def evidence_id(analysis_run_id, artifact_id, stage, unit_index, locator, ch):
    # Deliberately omit analysis_run_id so evidence IDs remain stable across reruns
    # when the artifact, stage, locator, index, and content are unchanged.
    return hashlib.sha1('|'.join(map(str,[artifact_id,stage,unit_index,locator,ch])).encode()).hexdigest()
def _unit(run,row,idx,typ,text,locator,heading=''):
    ch=content_hash(text); return {"evidence_unit_id":evidence_id(run,row['artifact_id'],row.get('stage'),idx,locator,ch),"analysis_run_id":run,"artifact_id":row['artifact_id'],"collection_id":row.get('collection_id'),"stage":row.get('stage'),"stage_order":row.get('stage_order'),"team_id":row.get('team_id'),"participant_id":row.get('participant_id'),"unit_index":str(idx),"unit_type":typ,"text":text,"normalized_text":normalize_text(text),"source_locator":locator,"section_heading":heading,"is_prompt":"false","is_response":"true","is_template_content":"false","content_hash":ch}
class _HTML(html.parser.HTMLParser):
    tags={'h1','h2','h3','h4','h5','h6','p','li','td','th'}
    def __init__(self): super().__init__(); self.cur=None; self.buf=[]; self.out=[]; self.i=0; self.heading=''
    def handle_starttag(self, tag, attrs):
        if tag in self.tags: self.cur=tag; self.buf=[]
    def handle_data(self, data):
        if self.cur: self.buf.append(data)
    def handle_endtag(self, tag):
        if tag==self.cur:
            text=' '.join(''.join(self.buf).split())
            if text:
                self.i+=1
                if tag.startswith('h'): self.heading=text
                self.out.append((tag,text,f"html:{self.i}:{tag}", self.heading if not tag.startswith('h') else text))
            self.cur=None; self.buf=[]


def _canonical_path(value: str | None) -> str:
    if not value:
        return ''
    raw = str(value).strip().replace('\\', '/')
    raw = raw.rstrip('/\\')
    if not raw:
        return ''
    try:
        raw = str(Path(raw).expanduser().resolve(strict=False)).replace('\\', '/')
    except Exception:
        raw = raw.replace('\\', '/')
    return raw.rstrip('/\\').casefold()

def _row_source_values(d: dict, *names: str) -> list[str]:
    vals=[]
    for name in names:
        v=d.get(name)
        if v and str(v) not in vals:
            vals.append(str(v))
    return vals

def _scope_rows(rows: list[dict], artifact_path: str | None, source_names: tuple[str, ...], label: str) -> tuple[list[dict], list[str]]:
    if not artifact_path:
        return rows, []
    target=_canonical_path(artifact_path)
    exact=[d for d in rows if any(_canonical_path(v)==target for v in _row_source_values(d, *source_names))]
    if exact:
        return exact, []
    distinct={}
    for d in rows:
        for v in _row_source_values(d, *source_names):
            cv=_canonical_path(v)
            if cv:
                distinct.setdefault(cv, v)
    if len(distinct)==1:
        only=next(iter(distinct.values()))
        return rows, [f"Could not exactly scope {label} evidence to artifact path {artifact_path}; using the only {label} source present: {only}."]
    if len(distinct)>1:
        return [], [f"Could not scope {label} evidence to artifact path {artifact_path}; multiple {label} sources are present."]
    return rows, []

def extract_file_units(path: str, run: str, row: dict) -> tuple[list[dict], list[str]]:
    p=Path(path); ext=(p.suffix or '').lower(); warnings=[]; units=[]
    try:
        if ext=='.txt':
            for n,line in enumerate(p.read_text(encoding='utf-8', errors='replace').splitlines(),1):
                txt=line.strip()
                if txt: units.append(_unit(run,row,len(units),'text_line',txt,f"line:{n}"))
        elif ext in {'.html','.htm'}:
            parser=_HTML(); parser.feed(p.read_text(encoding='utf-8', errors='replace'))
            for tag,text,loc,head in parser.out: units.append(_unit(run,row,len(units),f"html_{tag}",text,loc,head))
        elif ext=='.docx':
            try:
                from docx import Document
            except Exception:
                return [], ["python-docx is unavailable; DOCX evidence extraction skipped"]
            doc=Document(str(p))
            for i,para in enumerate(doc.paragraphs):
                txt=para.text.strip()
                if txt: units.append(_unit(run,row,len(units),'docx_paragraph',txt,f"paragraph:{i}"))
            for ti,t in enumerate(doc.tables):
                for ri,r in enumerate(t.rows):
                    for ci,c in enumerate(r.cells):
                        txt=c.text.strip()
                        if txt: units.append(_unit(run,row,len(units),'docx_table_cell',txt,f"table:{ti}:row:{ri}:cell:{ci}"))
        else: warnings.append(f"unsupported file extension for progression extraction: {ext}")
    except Exception as e: warnings.append(f"could not read artifact content: {e}")
    return units,warnings

def extract_tldraw_units(db_path: str, run: str, row: dict, artifact_path: str | None = None) -> tuple[list[dict], list[str]]:
    conn=connect(db_path); init_db(conn); units=[]; warnings=[]
    for table in ('parsed_tldraw_events','tldraw_events'):
        try: raw_rows=conn.execute(f"select * from {table}").fetchall()
        except Exception: continue
        rows=[dict(r) for r in raw_rows]
        scoped, warnings = _scope_rows(rows, artifact_path, ('origin','source_path'), 'TLDraw')
        for d in scoped:
            txt=(d.get('text_snippet') or d.get('text') or '').strip()
            if txt:
                source=d.get('origin') or d.get('source_path') or ''
                loc=f"table:{table}:source:{source}:object_id:{d.get('object_id','')}:event_key:{d.get('event_key','')}"
                units.append(_unit(run,row,len(units),'tldraw_text',txt,loc))
        break
    conn.close(); return units, warnings

def extract_transcript_units(db_path: str, run: str, row: dict, artifact_path: str | None = None) -> tuple[list[dict], list[str]]:
    conn=connect(db_path); init_db(conn); units=[]; warnings=[]
    try: raw_rows=conn.execute('select * from transcript_turn_events').fetchall()
    except Exception: raw_rows=[]
    rows=[dict(r) for r in raw_rows]
    scoped, warnings = _scope_rows(rows, artifact_path, ('source_path',), 'transcript')
    for d in scoped:
        txt=(d.get('text') or d.get('utterance') or '').strip()
        if txt:
            loc=f"source_path:{d.get('source_path','')}:turn_index:{d.get('turn_index','')}:start_seconds:{d.get('start_seconds','')}:end_seconds:{d.get('end_seconds','')}"
            units.append(_unit(run,row,len(units),'speech_turn',txt,loc))
    conn.close(); return units, warnings
